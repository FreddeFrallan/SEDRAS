"""
Reusable XLSX block helpers that can be composed by the rendering pipeline.

The blocks intentionally cover a range of spreadsheet techniques (merged cells,
conditional formatting, charts) so that the downstream PDFs converted from
these workbooks look varied and realistic.
"""

from __future__ import annotations

from collections import Counter
from typing import Any, Dict, Iterable, List, Sequence, Set, Tuple

from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

IndexedSample = Tuple[int, Dict[str, Any]]

_XLSX_BLOCKS: Dict[str, Any] = {}


def _register_block(func):
    _XLSX_BLOCKS[func.__name__] = func
    return func


def register_default_xlsx_blocks() -> List[str]:
    """
    Returns a sorted list of registered XLSX block names.
    """
    return sorted(_XLSX_BLOCKS.keys())


def get_registered_xlsx_blocks() -> Dict[str, Any]:
    """
    Returns a copy of the registered XLSX block callables.
    """
    return dict(_XLSX_BLOCKS)


def _normalize_indexed_samples(samples: Iterable[IndexedSample]) -> List[IndexedSample]:
    normalized: List[IndexedSample] = []
    for idx, sample in samples:
        if sample is None:
            continue
        normalized.append((idx, sample))
    return normalized


def _extract_property_strings(sample: Dict[str, Any]) -> list[str]:
    property_entries: list[str] = []
    label_meta = sample.get("instance_label_text") or {}
    property_meta = {}
    if isinstance(label_meta, dict):
        property_meta = label_meta.get("properties") or {}
        if not isinstance(property_meta, dict):
            property_meta = {}

    prop_labels = sample.get("property_labels") or {}
    for name, raw_value in prop_labels.items():
        display = property_meta.get(name)
        if display is None:
            display = "N/A" if raw_value is None else str(raw_value)
        property_entries.append(f"{name}: {display}")

    return property_entries


def _generate_label_string(
    sample: Dict[str, Any],
    label_mapping: Dict[int, str] | None = None,
) -> str:
    label_idx = sample.get("label")
    label_text = None

    if label_mapping and label_idx in label_mapping:
        label_text = label_mapping[label_idx]

    label_meta = sample.get("instance_label_text")
    print(f"Label meta: {label_meta}")
    # if not label_text and isinstance(label_meta, dict):
    label_text = label_meta["label"]

    label_display = "Unknown"
    if label_idx is not None:
        label_display = str(label_idx)
    if label_text:
        label_display = f"{label_display} ({label_text})"

    property_entries = _extract_property_strings(sample)
    if property_entries:
        label_display = f"{label_display} | " + "; ".join(property_entries)

    return label_display


def _auto_fit_columns(ws: Worksheet, freeze_header: bool = True) -> None:
    # Standard column width logic
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 60)

    # NEW: Keep the first row (headers) visible when scrolling down
    if freeze_header:
        ws.freeze_panes = "A2"


@_register_block
def render_placeholder_xlsx_block(*, worksheet: Worksheet, data: Sequence[Dict[str, Any]]) -> Set[int]:
    """
    Minimal placeholder renderer that writes column headers for quick smoke tests.
    """
    if not worksheet or not data:
        return set()
    headers = list(data[0].keys())
    worksheet.append(headers)
    for row in data:
        worksheet.append([row.get(h, "") for h in headers])
    _auto_fit_columns(worksheet)
    return set()


@_register_block
def render_cover_sheet(
        *,
        worksheet: Worksheet,
        title: str,
        subtitle: str | None = None,
        summary_items: Dict[str, str] | None = None,
) -> Set[int]:
    """
    Renders a high-impact cover page for the XLSX report.
    - Hides gridlines for a clean 'UI' look.
    - Uses merged cells and background fills for a professional header.
    - Formats metadata as a clean key-value card.
    """
    from openpyxl.styles import Border, Side

    # 1. Visual Cleanup: Hide gridlines and disable scrolling locks
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = None

    # 2. Main Title Header (Merged Span)
    # Spans across A1 to F3
    worksheet.merge_cells(start_row=1, start_column=1, end_row=3, end_column=6)
    title_cell = worksheet.cell(row=1, column=1, value=title.upper())
    title_cell.font = Font(bold=True, size=22, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill("solid", fgColor="2E4053")  # Deep Navy Blue

    # 3. Subtitle (Merged Span)
    worksheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=6)
    sub_text = subtitle or "Automated Data Analysis Report"
    sub_cell = worksheet.cell(row=4, column=1, value=sub_text)
    sub_cell.font = Font(italic=True, size=12, color="566573")
    sub_cell.alignment = Alignment(horizontal="center")

    # 4. Metadata Table ("Key Insights" Card)
    # We place this starting at row 6
    worksheet.append([])  # Spacer row
    start_row = 6
    headers = ["REPORT METRIC", "VALUE"]
    worksheet.append(headers)

    # Style the Metadata Headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="5DADE2")  # Light Blue
    border_style = Border(
        left=Side(style='thin', color='D5DBDB'),
        right=Side(style='thin', color='D5DBDB'),
        top=Side(style='thin', color='D5DBDB'),
        bottom=Side(style='thin', color='D5DBDB')
    )

    for col_idx, text in enumerate(headers, start=1):
        cell = worksheet.cell(row=start_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border_style

    # 5. Populate Summary Items
    items = summary_items or {}
    curr_row = start_row + 1
    for key, value in items.items():
        worksheet.cell(row=curr_row, column=1, value=key).font = Font(bold=True)
        worksheet.cell(row=curr_row, column=2, value=value)

        # Apply borders to the key-value cells
        worksheet.cell(row=curr_row, column=1).border = border_style
        worksheet.cell(row=curr_row, column=2).border = border_style

        # Zebra striping for metadata
        if curr_row % 2 == 0:
            fill = PatternFill("solid", fgColor="F4F6F7")
            worksheet.cell(row=curr_row, column=1).fill = fill
            worksheet.cell(row=curr_row, column=2).fill = fill

        curr_row += 1

    # Finalize column widths without adding freeze panes
    _auto_fit_columns(worksheet, freeze_header=False)

    return set()


@_register_block
def render_label_summary_sheet(
    *,
    worksheet: Worksheet,
    indexed_samples: Sequence[IndexedSample],
    label_mapping: Dict[int, str] | None = None,
) -> Set[int]:
    """
    Draws a matrix of label statistics with merged headers, data bars, and a small chart.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    used = {idx for idx, _ in samples}
    if not samples:
        return set()

    label_counts = Counter(sample.get("label") for _, sample in samples)
    total = sum(label_counts.values())
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    worksheet.cell(row=1, column=1, value="Label Distribution Overview").font = Font(bold=True, size=14)

    header = ["Label", "Name", "Count", "Percent", "Notes"]
    worksheet.append([])
    worksheet.append(header)
    header_row = worksheet.max_row
    for col in range(1, len(header) + 1):
        cell = worksheet.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F618D")
        cell.alignment = Alignment(horizontal="center")

    for label_value, count in sorted(label_counts.items()):
        percent = f"{(count / total) * 100:0.1f}%"
        label_name = (label_mapping or {}).get(label_value, f"Label {label_value}")
        worksheet.append([
            label_value,
            label_name,
            count,
            percent,
            "High volume" if count / total > 0.4 else "Balanced",
        ])

    data_bar_rule = DataBarRule(start_type="num", start_value=0, end_type="num", end_value=max(label_counts.values()))
    worksheet.conditional_formatting.add(f"C{header_row + 1}:C{worksheet.max_row}", data_bar_rule)

    chart = BarChart()
    chart.title = "Label Counts"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Label"
    data_ref = Reference(worksheet, min_col=3, min_row=header_row, max_row=worksheet.max_row)
    cats_ref = Reference(worksheet, min_col=1, min_row=header_row + 1, max_row=worksheet.max_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    worksheet.add_chart(chart, f"H{header_row}")

    _auto_fit_columns(worksheet)
    return used


@_register_block
def render_featured_samples_grid(
    *,
    worksheet: Worksheet,
    indexed_samples: Sequence[IndexedSample],
    columns: int = 2,
    wrap_text: bool = True,
) -> Set[int]:
    """
    Places sample snippets into a storyboard grid, merging cells for the description rows.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    used = set()
    if not samples:
        return used

    worksheet.append(["Sample Metadata"] * columns)
    worksheet.append([])
    row_pointer = worksheet.max_row + 1

    col_width = 40
    for col in range(1, columns * 2 + 1):
        worksheet.column_dimensions[get_column_letter(col)].width = col_width if col % 2 == 0 else 15

    for chunk_start in range(0, len(samples), columns):
        chunk = samples[chunk_start:chunk_start + columns]
        max_rows = 4
        for idx, sample in chunk:
            used.add(idx)
        for col_offset, (sample_idx, sample) in enumerate(chunk):
            col_base = col_offset * 2 + 1
            worksheet.cell(row=row_pointer, column=col_base, value=f"ID {sample_idx}")
            worksheet.cell(row=row_pointer, column=col_base + 1, value=_generate_label_string(sample))
            worksheet.cell(row=row_pointer + 1, column=col_base, value="Label text")
            label_text = (sample.get("instance_label_text") or {}).get("label") if isinstance(sample.get("instance_label_text"), dict) else ""
            worksheet.cell(row=row_pointer + 1, column=col_base + 1, value=label_text or "n/a")

            worksheet.merge_cells(
                start_row=row_pointer + 2,
                end_row=row_pointer + 1 + max_rows,
                start_column=col_base,
                end_column=col_base + 1,
            )
            snippet_cell = worksheet.cell(row=row_pointer + 2, column=col_base)
            snippet_cell.value = (sample.get("instance_text") or "")[:600]
            snippet_cell.alignment = Alignment(wrap_text=wrap_text, vertical="top")
            snippet_cell.fill = PatternFill("solid", fgColor="FCF3CF")

        row_pointer += max_rows + 2
        worksheet.append([])

    return used


@_register_block
def render_raw_samples_sheet(
    *,
    worksheet: Worksheet,
    indexed_samples: Sequence[IndexedSample],
    label_mapping: Dict[int, str] | None = None,
) -> Set[int]:
    """
    Renders the full sample roster with colored headers and wrapped text cells.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    used = {idx for idx, _ in samples}
    if not samples:
        return set()

    headers = ["Sample ID", "Label", "Label Name", "Instance Text"]
    worksheet.append(headers)
    header_row = worksheet.max_row
    for col in range(1, len(headers) + 1):
        cell = worksheet.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor="1A5276")

    for idx, sample in samples:
        label = sample.get("label")
        label_display = _generate_label_string(sample, label_mapping)
        label_name = (label_mapping or {}).get(label, f"Label {label}")
        text = sample.get("instance_text") or ""
        worksheet.append([idx, label_display, label_name, text])
        row_idx = worksheet.max_row
        for col in range(1, len(headers) + 1):
            cell = worksheet.cell(row=row_idx, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F9EBEA")
            else:
                cell.fill = PatternFill("solid", fgColor="FEF5E7")

    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        worksheet.column_dimensions[letter].width = 20 if col < 4 else 60

    return used


@_register_block
def render_label_ribbon_matrix(
    *,
    worksheet: Worksheet,
    indexed_samples: Sequence[IndexedSample],
    palette: Sequence[str] | None = None,
) -> Set[int]:
    """
    Creates a color-coded matrix with rotated label headers to emphasize label differences.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    used = {idx for idx, _ in samples}
    if not samples:
        return set()

    palette = palette or ["1ABC9C", "F5B041", "5DADE2", "AF7AC5", "E74C3C"]
    worksheet.append(["Sample ID", "Label", "Excerpt"])
    header_row = worksheet.max_row
    for col in range(1, 4):
        cell = worksheet.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="34495E")
        cell.alignment = Alignment(horizontal="center")

    for idx, sample in samples:
        text = (sample.get("instance_text") or "")[:400]
        label = sample.get("label") or 0
        label_display = _generate_label_string(sample)
        worksheet.append([idx, label_display, text])
        row = worksheet.max_row
        label_cell = worksheet.cell(row=row, column=2)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", textRotation=90)
        fill_color = palette[label % len(palette)]
        label_cell.fill = PatternFill("solid", fgColor=fill_color)
        label_cell.font = Font(bold=True, color="FFFFFF")
        text_cell = worksheet.cell(row=row, column=3)
        text_cell.alignment = Alignment(wrap_text=True, vertical="top")
        text_cell.fill = PatternFill("solid", fgColor="FDFEFE")

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 6
    worksheet.column_dimensions["C"].width = 70
    return used


@_register_block
def render_dual_column_landscape(
    *,
    worksheet: Worksheet,
    indexed_samples: Sequence[IndexedSample],
    columns: int = 2,
) -> Set[int]:
    """
    Splits the sheet into multiple card-style columns with merged cells and accent colors.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    used = set()
    if not samples:
        return used

    card_width = 30
    for col in range(1, columns * 3 + 1):
        worksheet.column_dimensions[get_column_letter(col)].width = card_width if col % 3 == 0 else 12

    row_pointer = 1
    palette = ["EBF5FB", "FDEDEC", "F9EBEA", "EAFAF1"]

    for chunk_start in range(0, len(samples), columns):
        chunk = samples[chunk_start:chunk_start + columns]
        height = 6
        for col_offset, (idx, sample) in enumerate(chunk):
            used.add(idx)
            base_col = col_offset * 3 + 1
            worksheet.merge_cells(
                start_row=row_pointer,
                end_row=row_pointer + height,
                start_column=base_col,
                end_column=base_col + 2,
            )
            cell = worksheet.cell(row=row_pointer, column=base_col)
            text = sample.get("instance_text") or ""
            label_display = _generate_label_string(sample)
            cell.value = f"Sample {idx} | {label_display}\n{text}"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = PatternFill("solid", fgColor=palette[col_offset % len(palette)])
            cell.font = Font(size=11)
        row_pointer += height + 2

    return used


@_register_block
def render_split_text_table(
    *,
    worksheet: Worksheet,
    indexed_samples: Sequence[IndexedSample],
    max_tokens: int = 200,
) -> Set[int]:
    """
    Splits each sample's text into multiple rows, placing each chunk in its own cell.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    used = set()
    if not samples:
        return used

    header = ["Sample ID", "Label", "Chunk #", "Text Slice"]
    worksheet.append(header)
    header_row = worksheet.max_row
    for col in range(1, len(header) + 1):
        cell = worksheet.cell(row=header_row, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor="1B4F72")

    for idx, sample in samples:
        text = sample.get("instance_text") or ""
        label_display = _generate_label_string(sample)
        chunks = [text[i:i + max_tokens] for i in range(0, len(text), max_tokens)] or [""]
        for chunk_idx, chunk in enumerate(chunks, start=1):
            worksheet.append([idx if chunk_idx == 1 else "", label_display if chunk_idx == 1 else "", chunk_idx, chunk])
            row_idx = worksheet.max_row
            slice_cell = worksheet.cell(row=row_idx, column=4)
            slice_cell.alignment = Alignment(wrap_text=True, vertical="top")
            slice_cell.fill = PatternFill("solid", fgColor="FAF9F6")
            used.add(idx)

    worksheet.column_dimensions[get_column_letter(1)].width = 12
    worksheet.column_dimensions[get_column_letter(2)].width = 10
    worksheet.column_dimensions[get_column_letter(3)].width = 10
    worksheet.column_dimensions[get_column_letter(4)].width = 80
    return used


@_register_block
def render_id_label_reference(
    *,
    worksheet: Worksheet,
    indexed_samples: Sequence[IndexedSample],
    label_mapping: Dict[int, str] | None = None,
) -> Set[int]:
    """
    First creates an ID/label lookup table, then a separate ID/text table on the same sheet.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    used = {idx for idx, _ in samples}
    if not samples:
        return set()

    # Section 1: ID -> Label map
    worksheet.append(["ID ↔ Label Lookup"])
    worksheet.cell(row=worksheet.max_row, column=1).font = Font(bold=True, size=14)
    worksheet.append(["Sample ID", "Label", "Label Name"])
    header_row = worksheet.max_row
    for col in range(1, 4):
        cell = worksheet.cell(row=header_row, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor="1F618D")

    for idx, sample in samples:
        label = sample.get("label")
        label_display = _generate_label_string(sample, label_mapping)
        label_name = (label_mapping or {}).get(label, f"Label {label}")
        worksheet.append([idx, label_display, label_name])

    worksheet.append([])
    worksheet.append(["ID ↔ Text Reference"])
    worksheet.cell(row=worksheet.max_row, column=1).font = Font(bold=True, size=14)
    worksheet.append(["Sample ID", "Text"])
    header_row = worksheet.max_row
    for col in range(1, 3):
        cell = worksheet.cell(row=header_row, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor="117A65")

    for idx, sample in samples:
        text = sample.get("instance_text") or ""
        worksheet.append([idx, text])
        row_idx = worksheet.max_row
        text_cell = worksheet.cell(row=row_idx, column=2)
        text_cell.alignment = Alignment(wrap_text=True, vertical="top")

    worksheet.column_dimensions["A"].width = 15
    worksheet.column_dimensions["B"].width = 70
    worksheet.column_dimensions["C"].width = 30
    return used


@_register_block
def render_property_heatmap(
    *,
    worksheet: Worksheet,
    indexed_samples: Sequence[IndexedSample],
    property_keys: Sequence[str] | None = None,
) -> Set[int]:
    """
    Builds a numeric matrix highlighting per-sample property scores with a color scale.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    used = {idx for idx, _ in samples}
    if not samples:
        return set()

    worksheet.append(["Sample ID"] + list(property_keys or ["property_a", "property_b", "property_c"]))
    header_row = worksheet.max_row
    for col in range(1, worksheet.max_column + 1):
        cell = worksheet.cell(row=header_row, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D4E6F1")

    for idx, sample in samples:
        properties = sample.get("property_scores") or sample.get("property_labels") or {}
        row = [idx]
        for key in property_keys or ["property_a", "property_b", "property_c"]:
            value = properties.get(key)
            if value is None:
                value = properties.get(key.replace("property_", ""), "")
            row.append(value if isinstance(value, (int, float)) else 0)
        worksheet.append(row)

    worksheet.conditional_formatting.add(
        f"B{header_row + 1}:{get_column_letter(worksheet.max_column)}{worksheet.max_row}",
        ColorScaleRule(
            start_type="percentile", start_value=10, start_color="F2F3F4",
            mid_type="percentile", mid_value=50, mid_color="85C1E9",
            end_type="percentile", end_value=90, end_color="1B4F72",
        ),
    )

    _auto_fit_columns(worksheet)
    return used


@_register_block
def render_running_totals_sheet(
        *,
        worksheet: Worksheet,
        indexed_samples: Sequence[IndexedSample],
        window: int = 5,
) -> Set[int]:
    """
    Calculates rolling window statistics across the provided sample set.
    Includes explicit colors for DataBars to prevent openpyxl NoneType errors.
    """
    samples = _normalize_indexed_samples(indexed_samples)
    if not samples:
        return set()

    # 1. Sheet Setup: Hide gridlines for a cleaner look
    worksheet.sheet_view.showGridLines = False
    used = {idx for idx, _ in samples}

    # 2. Header Row
    headers = ["Sample Index", "Text Length", "Rolling Avg"]
    worksheet.append(headers)
    header_row = worksheet.max_row

    for cell in worksheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4053")  # Professional Navy
        cell.alignment = Alignment(horizontal="center")

    # 3. Calculate and Append Data
    lengths: List[int] = []
    for position, (idx, sample) in enumerate(samples, start=1):
        text = sample.get("instance_text") or ""
        lengths.append(len(text))

        # Maintain the rolling window
        if len(lengths) > window:
            lengths.pop(0)

        rolling_avg = sum(lengths) / len(lengths)

        worksheet.append([idx, len(text), round(rolling_avg, 1)])

        # Zebra striping for readability
        row_idx = worksheet.max_row
        if row_idx % 2 == 0:
            for col in range(1, 4):
                worksheet.cell(row=row_idx, column=col).fill = PatternFill("solid", fgColor="F4F6F7")

    # 4. Apply Conditional Formatting (The Fix)
    # We must provide an explicit hex color string to avoid the TypeError
    data_bar_rule = DataBarRule(
        start_type="num",
        start_value=0,
        end_type="max",
        color="5DADE2",  # Sky Blue hex
        showValue=True
    )

    # Apply to the 'Rolling Avg' column (Column C)
    worksheet.conditional_formatting.add(
        f"C{header_row + 1}:C{worksheet.max_row}",
        data_bar_rule
    )

    # 5. Column Formatting & Freeze Panes
    # We lock the top row so the user doesn't lose the headers while scrolling
    _auto_fit_columns(worksheet, freeze_header=True)

    return used
