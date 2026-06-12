#!/usr/bin/env python3
from __future__ import annotations

import inspect
import json
import os
import random
import re
from datetime import datetime
from typing import Any, Callable, Dict, List, Optional, Sequence, Set, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font

from data_management.rendering.errors import InvalidTemplateError
from data_management.rendering.xlsx.generate_xlsx_layout import XlsxLayoutGenerator
from data_management.rendering.xlsx.code_blocks.general_xlsx_blocks import get_registered_xlsx_blocks
from data_management.rendering.trace_logger import RenderingTraceLogger
from inference.model_wrappers.llm_wrapper import get_llm_wrapper, LLMModel, LLMWrapper


# --- ROBUST UTILS ---

def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("_") or "unnamed"


def _clean_for_excel(v: Any) -> str:
    """Removes newlines and illegal characters that break XLSX cell rendering."""
    if v is None: return ""
    s = str(v).replace('\n', ' ').replace('\r', ' ')
    return "".join(c for c in s if c.isprintable())


def _sanitize_xlsx_sheet_title(title: str) -> str:
    """
    Cleans a string to be a valid Excel sheet title.
    - Removes forbidden characters: \ / * ? : [ ]
    - Limits length to 31 characters.
    - Ensures it is not empty.
    """
    # 1. Replace forbidden characters with a space or underscore
    # Forbidden: \ / * ? : [ ]
    clean_title = re.sub(r'[\\/*?:\[\]]', ' ', title)

    # 2. Excel sheet names cannot exceed 31 characters
    clean_title = clean_title.strip()[:31]

    # 3. Fallback if the title becomes empty after cleaning
    return clean_title or "Sheet"


def _find_instance_json(dataset_root: str, instance_name: str, level: Optional[str]) -> Tuple[str, Dict[str, Any]]:
    instances_root = os.path.join(dataset_root, "instances")
    for root, _, files in os.walk(instances_root):
        for f in files:
            if f in (f"{instance_name}.json", f"{_slug(instance_name)}.json"):
                path = os.path.join(root, f)
                with open(path, "r", encoding="utf-8") as f_in:
                    return path, json.load(f_in)
    raise FileNotFoundError(f"Instance '{instance_name}' not found in {instances_root}")


class _SafeDict(dict):
    def __missing__(self, key): return f"{{{key}}}"


def _sanitize_template(template: Any) -> str:
    """Strips Python logic like .split() which the LLM keeps trying to use."""
    if not isinstance(template, str): return str(template or "")
    return re.sub(r"\{(\w+)(\.[^}]+)\}", r"{\1}", template)


def _render_template(template: Any, context: Dict[str, Any]) -> str:
    cleaned = _sanitize_template(template)
    safe_context = {k: _clean_for_excel(v) for k, v in context.items()}
    try:
        return cleaned.format_map(_SafeDict(**safe_context))
    except Exception:
        res = cleaned
        for k, v in safe_context.items():
            res = res.replace(f"{{{k}}}", v)
        return res


def _extract_block_text(block: Dict[str, Any]) -> str:
    """Detects which key the LLM used for static text content."""
    # SOTA models sometimes hallucinate 'prompt' or 'content' from other tasks
    for key in ["text", "prompt", "title", "content"]:
        val = block.get(key)
        if val and isinstance(val, str) and not val.endswith("()"):
            return val
    return ""


def _unique_sheet_title(base_title: str, used_titles: Set[str]) -> str:
    sanitized = (base_title or "Sheet").strip() or "Sheet"
    title = sanitized[:31]
    if title not in used_titles:
        used_titles.add(title)
        return title
    for counter in range(2, 100):
        suffix = f"_{counter}"
        trimmed = title[:31 - len(suffix)] + suffix
        if trimmed not in used_titles:
            used_titles.add(trimmed)
            return trimmed
    fallback = f"Sheet_{len(used_titles) + 1}"
    used_titles.add(fallback)
    return fallback


def _build_indexed_subset(samples: Sequence[Dict[str, Any]], indices: Sequence[int]) -> List[Tuple[int, Dict[str, Any]]]:
    return [(idx, samples[idx]) for idx in indices if 0 <= idx < len(samples)]


def _select_table_block_renderer() -> List[Tuple[str, Callable[..., Set[int]], inspect.Signature]]:
    candidates: List[Tuple[str, Callable[..., Set[int]], inspect.Signature]] = []
    for name, func in get_registered_xlsx_blocks().items():
        if name == "render_cover_sheet":
            continue
        signature = inspect.signature(func)
        params = signature.parameters
        if "worksheet" not in params:
            continue
        if "indexed_samples" in params or "data" in params:
            candidates.append((name, func, signature))
    if not candidates:
        raise InvalidTemplateError("No registered XLSX table block renderers are available.")
    return candidates


# --- MAIN RENDERER ---

def render_dataset_instance_to_xlsx(
        dataset_root: str,
        instance_name: str,
        *,
        level: Optional[str] = None,
        instance_title: Optional[str] = None,
        out_xlsx: Optional[str] = None,
        llm: Optional[LLMWrapper] = None,
        llm_model: Optional[str] = None,
        style_hint: Optional[str] = None,
        prompt_md_path: Optional[str] = None,
        instance_data: Optional[Dict[str, Any]] = None,
        samples: Optional[List[Dict[str, Any]]] = None,
        error_feedback: Optional[str] = None,
        previous_response: Optional[Any] = None,
        trace_logger: Optional[RenderingTraceLogger] = None,
) -> str:
    dataset_root = os.path.abspath(dataset_root)
    inst = instance_data
    if inst is None:
        _, inst = _find_instance_json(dataset_root, instance_name, level)
    samples = samples if samples is not None else inst.get("dataset", [])

    tm = inst.get("text_mapping", {})
    ol = tm.get("output_labels", {}) if isinstance(tm, dict) else {}
    label_names = {int(k): str(v) for k, v in ol.items() if str(k).isdigit()}

    title = instance_title or inst.get("name") or instance_name
    unused_indices = set(range(len(samples)))
    covered_indices: Set[int] = set()

    if trace_logger:
        trace_logger.log_stage("xlsx_render_start", samples=len(samples), instance_title=title)

    effective_llm = llm or get_llm_wrapper(llm_model or LLMModel.GEMINI_3_PRO)
    gen = XlsxLayoutGenerator(prompt_md=prompt_md_path)
    layout = gen.generate_layout(
        llm=effective_llm,
        instance_title=title,
        samples=samples,
        style_hint=style_hint,
        error_feedback=error_feedback,
        previous_response=previous_response,
        trace_logger=trace_logger,
    )

    renders_dir = os.path.join(dataset_root, "renders", "interim_xlsx")
    if level: renders_dir = os.path.join(renders_dir, level)
    os.makedirs(renders_dir, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_fn = f"{_slug(title)}_{timestamp}"
    layout_json_path = os.path.join(renders_dir, f"{base_fn}_layout.json")
    with open(layout_json_path, "w", encoding="utf-8") as f:
        json.dump(layout, f, indent=2, ensure_ascii=False)

    wb = Workbook()
    ws = wb.active
    ws.title = (layout.get("sheet_name") or "Report")[:31]
    used_sheet_titles = {ws.title}
    table_block_renderers = _select_table_block_renderer()

    MANDATORY = [
        {"header": "sample_id", "value": "{sample_id}"},
        {"header": "text", "value": "{instance_text}"},
        {"header": "label", "value": "{label_details}"},
    ]

    for b_idx, block in enumerate(layout.get("blocks", [])):
        if isinstance(block, str):
            raise InvalidTemplateError(f"Block {b_idx} is a string ('{block}'). Expected a JSON object.")

        b_type = str(block.get("type", "table_block")).lower()
        # print(f"Processing block {b_idx}: type='{b_type}' title='{block.get('title') or 'No Title'}'")

        # --- PATH A: Content Blocks (Headers, Paragraphs) ---
        if b_type in ["header_block", "text_block", "title_block", "intro_paragraph"]:
            text = _extract_block_text(block)
            if text:
                ws.append([_clean_for_excel(text)])
                # Style headers to be Bold
                if "header" in b_type or "title" in b_type:
                    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
                ws.append([])  # Spacer row
            # print(f"⚠️ Skipping empty content block {b_idx} ('{block.get('title') or 'No Title'}')")
            continue

        # --- PATH B: Data Table Blocks ---
        req = block.get("sample_indices", [])
        indices = sorted(list(unused_indices)) if req == "all_remaining" else [i for i in req if i in unused_indices]

        # If it's labeled a table but has no indices, treat it as text to avoid silent failure
        if not indices:
            fallback = _extract_block_text(block)
            if fallback:
                ws.append([_clean_for_excel(fallback)])
                ws.append([])
            # print(f"⚠️ Skipping empty table block {b_idx} ('{block.get('title') or 'No Title'}')")
            continue

        unused_indices.difference_update(indices)
        block_title = block.get("title") or block.get("text") or f"Table {b_idx + 1}"
        # First, sanitize the title string
        block_title = _sanitize_xlsx_sheet_title(str(block_title))

        sheet_title = _unique_sheet_title(str(block_title), used_sheet_titles)
        table_ws = wb.create_sheet(title=sheet_title)

        renderer_name, renderer_fn, renderer_sig = random.choice(table_block_renderers)
        subset = _build_indexed_subset(samples, indices)
        # print(f"Selected renderer '{renderer_name}' for block {b_idx}, covering {len(subset)} samples.")

        render_kwargs: Dict[str, Any] = {"worksheet": table_ws}
        if "indexed_samples" in renderer_sig.parameters:
            render_kwargs["indexed_samples"] = subset
        elif "data" in renderer_sig.parameters:
            render_kwargs["data"] = [sample for _, sample in subset]
        if "label_mapping" in renderer_sig.parameters:
            render_kwargs["label_mapping"] = label_names

        # print(f"Rendering table block {b_idx} using '{renderer_name}' on sheet '{sheet_title}'")
        used = renderer_fn(**render_kwargs)
        if not used:
            used = set(indices)
        covered_indices.update(used)

    # 4. Mandatory Coverage (Appendix)
    if unused_indices:
        # print(f"⚠️ Appending {len(unused_indices)} unassigned samples to the end of the XLSX.")
        ws.append(["AUTOMATIC APPENDIX: REMAINING SAMPLES"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="FF0000")
        ws.append([c["header"] for c in MANDATORY])
        for s_idx in sorted(list(unused_indices)):
            sample = samples[s_idx]
            ctx = {"sample_id": str(sample.get("sample_id") or s_idx),
                   "instance_text": sample.get("instance_text") or "", "label_details": str(sample.get("label"))}
            ws.append([_render_template(c["value"], ctx) for c in MANDATORY])
            covered_indices.add(s_idx)

    # 5. Final Validation & Formatting
    expected_indices = set(range(len(samples)))
    if covered_indices != expected_indices:
        missing = sorted(expected_indices - covered_indices)
        raise InvalidTemplateError(f"Missing sample indices: {missing}")

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 70)

    path = out_xlsx or os.path.join(renders_dir, f"{base_fn}.xlsx")
    wb.save(path)
    return path
